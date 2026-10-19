"""Importa os lançamentos da planilha CONSIPA (aba CentralReal) para o Financeiro.

Na CONSIPA cada PARCELA é uma linha solta. Aqui as linhas da mesma compra (mesma data,
classificação, fornecedor, descrição, nota e valor) viram UM lançamento com N parcelas.
Decisões aplicadas (16/10/2026): "Pecuária de leite" = fazenda Castro; agricultura -> Sede;
sem categoria vai pra fila de revisão; 3 categorias no galho errado são corrigidas.
Pode rodar de novo: o que já foi importado (chave_origem) é pulado, nada é apagado."""
from __future__ import annotations

import hashlib
import io
import warnings
from collections import defaultdict
from datetime import date, datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models_financeiro import FinFornecedor, FinLancamento, FinParcela, FinPlano
from .fin_lancamentos import normalizar
from .fin_plano import SETORES_CONSIPA, ErroFinanceiro, garantir_cadastros, limpar_nome

ABA = "CentralReal"
# Categoria cujo código é de outra central -> central certa (os 3 casos achados na planilha).
CORRECOES_GRUPO = {
    "1.04.02": "1.04. MATERIAIS DE CONSUMO, OPERAÇÃO E MANUTENÇÃO",
    "1.13.04": "1.13. CAPTAÇÃO DE CRÉDITO",
    "2.02.04": "2.02. VENDA DE ANIMAIS",
}
# Colunas digitadas da tabela tblCentralReal (a partir da coluna B).
COLUNAS = ["data", "data_fluxo", "status", "setor", "es", "central", "categoria", "item", "subitem",
           "nf", "fornecedor", "descricao", "n_parcelas", "qtde", "valor_unit", "arroba", "unidade",
           "rendimento", "desconto", "obs", "conta", "fazenda", "valor"]


def _vazio(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "0")) or v == 0


def _txt(v):
    return None if _vazio(v) else str(v).strip()


def _dt(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def ler_planilha(conteudo: bytes) -> list[dict]:
    """Lê a aba CentralReal (só as colunas digitadas) — modo leitura rápida."""
    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise ErroFinanceiro("Não consegui abrir o arquivo. Envie a planilha CONSIPA (.xlsm ou .xlsx).")
    if ABA not in wb.sheetnames:
        raise ErroFinanceiro("A planilha não tem a aba CentralReal — é mesmo a CONSIPA?")
    ws = wb[ABA]
    linhas = []
    cabecalho_ok = False
    for r in ws.iter_rows(min_row=1, min_col=2, max_col=24, values_only=True):
        if not cabecalho_ok:
            cabecalho_ok = str(r[0] or "").strip().upper().startswith("DATA CONT")
            continue
        if r[0] is None and r[4] is None:
            continue
        linhas.append(dict(zip(COLUNAS, r)))
    wb.close()
    if not cabecalho_ok:
        raise ErroFinanceiro("Não achei o cabeçalho da tabela na aba CentralReal")
    return linhas


# ---------------------------------------------------------------- Planilha antiga (anos anteriores)

# Nome da coluna (como está no cabeçalho, sem acento/maiúscula) -> campo do importador.
CABECALHOS = {
    "data contabil": "data", "data fluxo": "data_fluxo", "status": "status", "setor": "setor",
    "entradas/saidas": "es", "central de custo": "central", "categoria": "categoria", "item": "item",
    "subitem": "subitem", "no nota fiscal": "nf", "n nota fiscal": "nf", "fornecedor": "fornecedor",
    "descricao": "descricao", "no de parcelas": "n_parcelas", "n de parcelas": "n_parcelas",
    "quantidade": "qtde", "valor unitario": "valor_unit", "@/kg/l/saca": "arroba", "unidade": "unidade",
    "rendimento": "rendimento", "desconto / agio": "desconto", "observacao": "obs", "conta": "conta",
    "fazenda": "fazenda", "valor total": "valor",
}


def _sem_acento(t: str) -> str:
    import unicodedata
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(t.replace("º", "o").split())


def _num(v, padrao=0.0) -> float:
    return float(v) if isinstance(v, (int, float)) else padrao


def valor_pela_formula(r: dict) -> float:
    """Mesma conta da CONSIPA: qtde × unitário × (kg/@, quando houver) × (1 − desconto) ÷ parcelas.
    Usada quando o "valor total" da planilha está quebrado (#REF!)."""
    q = _num(r.get("qtde"), 1.0) or 1.0
    n = _num(r.get("n_parcelas"), 1.0) or 1.0
    mult = _num(r.get("arroba")) or 1.0
    d = _num(r.get("desconto"))
    bruto = q * _num(r.get("valor_unit")) * mult
    return round(((bruto - d) if d >= 1 else bruto * (1 - d)) / n, 2)


def ler_planilha_antiga(conteudo: bytes) -> list[dict]:
    """Lê qualquer aba com o cabeçalho da CONSIPA ("DATA CONTÁBIL" ...), achando as colunas
    pelo nome (as versões antigas não têm a coluna FAZENDA). Valor quebrado é recalculado."""
    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise ErroFinanceiro("Não consegui abrir a planilha antiga (.xlsx ou .xlsm).")
    linhas = []
    for ws in wb.worksheets:
        mapa = None
        for r in ws.iter_rows(values_only=True):
            if mapa is None:
                nomes = [_sem_acento(x) for x in r[:60]]
                if any(n.startswith("data contabil") for n in nomes):
                    mapa = {i: CABECALHOS[n] for i, n in enumerate(nomes) if n in CABECALHOS}
                continue
            if not any(r):
                continue
            d = {campo: (r[i] if i < len(r) else None) for i, campo in mapa.items()}
            d["_aba"] = ws.title
            if not isinstance(d.get("valor"), (int, float)) or not d.get("valor"):
                d["valor"] = valor_pela_formula(d)
            linhas.append(d)
    wb.close()
    if not linhas:
        raise ErroFinanceiro("Não achei nenhuma tabela com o cabeçalho da CONSIPA na planilha antiga.")
    return linhas


def _assinatura(r: dict) -> tuple:
    """Identifica o mesmo lançamento em planilhas diferentes (data, setor, E/S, categoria,
    fornecedor e valor) — usado pra não importar duas vezes."""
    return (_dt(r.get("data")), str(r.get("setor") or "").strip().upper(),
            str(r.get("es") or "").strip().upper(), (str(r.get("categoria") or "").strip().upper())[:10],
            str(r.get("fornecedor") or "").strip().upper(), round(_num(r.get("valor")), 0))


def combinar(atual: list[dict], antiga: list[dict], corte: date | None = None) -> tuple[list[dict], dict]:
    """Junta a CONSIPA atual com a planilha antiga.

    Da antiga entra só o que é ANTES do corte (padrão: primeiro mês em que a CONSIPA tem
    lançamentos de verdade) e que não está na atual. Do corte em diante vale a CONSIPA
    (é a versão mais nova e corrigida)."""
    validas = [r for r in atual if _dt(r.get("data")) and str(r.get("es") or "").strip().upper() in ("1. DESPESA", "2. RECEITA")]
    if corte is None:
        # Corte = primeiro mês em que a CONSIPA já tem pelo menos 80% das linhas que a
        # planilha antiga tem naquele mês (a CONSIPA "assumiu" o controle). Antes disso, a
        # antiga é a fonte; os poucos lançamentos que estão nas duas não se repetem.
        atual_mes = defaultdict(int)
        for r in validas:
            atual_mes[_dt(r["data"]).replace(day=1)] += 1
        antiga_mes = defaultdict(lambda: defaultdict(int))  # mês -> aba -> linhas
        for r in antiga:
            d = _dt(r.get("data"))
            if d and str(r.get("es") or "").strip().upper() in ("1. DESPESA", "2. RECEITA"):
                antiga_mes[d.replace(day=1)][r.get("_aba")] += 1
        corte = date.max
        for m in sorted(set(atual_mes) | set(antiga_mes)):
            maior_aba = max(antiga_mes[m].values(), default=0)
            if atual_mes[m] and atual_mes[m] >= 0.8 * maior_aba:
                corte = m
                break
    # Cada aba da planilha antiga é comparada com a CONSIPA E com as abas anteriores (a
    # mesma linha pode estar copiada em duas abas). Multiconjunto: duas compras iguais no
    # mesmo dia, na mesma aba, continuam sendo duas.
    from collections import Counter
    base = Counter(_assinatura(r) for r in atual)
    extras, repetidas, depois = [], 0, 0
    abas = list(dict.fromkeys(r.get("_aba") for r in antiga))
    for aba in abas:
        disponiveis = base.copy()
        da_aba = Counter()
        for r in (x for x in antiga if x.get("_aba") == aba):
            d = _dt(r.get("data"))
            if d is None or str(r.get("es") or "").strip().upper() not in ("1. DESPESA", "2. RECEITA"):
                continue
            if d >= corte:
                depois += 1
                continue
            a = _assinatura(r)
            da_aba[a] += 1
            if disponiveis[a] > 0:
                disponiveis[a] -= 1
                repetidas += 1
                continue
            extras.append(r)
        base = base | da_aba  # a próxima aba também não repete o que esta já trouxe
    return atual + extras, {"corte": corte.isoformat() if corte != date.max else None,
                            "antiga_incluidas": len(extras), "antiga_repetidas": repetidas,
                            "antiga_depois_do_corte": depois}


def importar(db: Session, conteudo: bytes, usuario: str, antiga: bytes | None = None) -> dict:
    """Importa a CONSIPA; com `antiga`, inclui também os anos anteriores da planilha antiga."""
    garantir_cadastros(db)
    linhas = ler_planilha(conteudo)
    info_antiga = None
    if antiga:
        linhas, info_antiga = combinar(linhas, ler_planilha_antiga(antiga))

    # Agrupa as parcelas da mesma compra.
    grupos: dict[tuple, list[dict]] = defaultdict(list)
    ignoradas = 0
    for r in linhas:
        es = str(r["es"] or "").strip().upper()
        if es not in ("1. DESPESA", "2. RECEITA") or not _dt(r["data"]) or _vazio(r["valor"]):
            ignoradas += 1  # inventário, configurações, linha sem valor
            continue
        chave = tuple(str(r[c]) for c in ("data", "setor", "es", "central", "categoria", "item", "nf",
                                          "fornecedor", "descricao", "n_parcelas", "qtde", "valor_unit", "conta"))
        grupos[chave].append(r)

    ja_importadas = set(db.scalars(select(FinLancamento.chave_origem).where(FinLancamento.origem == "consipa")))
    fornecedores = {}
    for f in db.scalars(select(FinFornecedor)):
        for n in [f.nome] + (f.apelidos or "").split(";"):
            if n.strip():
                fornecedores.setdefault(normalizar(n), f)
    plano = {(p.tipo, p.setor, p.grupo, p.categoria) for p in db.scalars(select(FinPlano))}

    criados = pulados = revisar = parcelas_total = 0
    for chave, rows in grupos.items():
        rows.sort(key=lambda x: _dt(x["data_fluxo"]) or _dt(x["data"]))
        n = max(1, int(float(rows[0]["n_parcelas"] or 1)))
        # Duas compras idênticas no mesmo dia: quebra em blocos do tamanho do parcelamento.
        for bloco_i in range(0, len(rows), n):
            bloco = rows[bloco_i:bloco_i + n]
            origem = hashlib.sha1(repr((chave, bloco_i)).encode()).hexdigest()
            if origem in ja_importadas:
                pulados += 1
                continue
            lanc = _montar(bloco, usuario, origem, fornecedores, plano, db)
            db.add(lanc)
            criados += 1
            parcelas_total += len(lanc.parcelas)
            revisar += lanc.revisar
    db.flush()
    _aprender_padroes(db)
    db.commit()
    r = {"lancamentos": criados, "parcelas": parcelas_total, "ja_existiam": pulados,
         "a_revisar": revisar, "linhas_ignoradas": ignoradas}
    if info_antiga:
        r["planilha_antiga"] = info_antiga
    return r


def _montar(bloco, usuario, origem, fornecedores, plano, db) -> FinLancamento:
    r = bloco[0]
    tipo = "despesa" if str(r["es"]).strip().upper().startswith("1") else "receita"
    setor_original = str(r["setor"] or "").strip().upper()
    setor = SETORES_CONSIPA.get(setor_original, (r["setor"] or "").strip().title() or None)
    obs = _txt(r["obs"])
    if obs and obs.replace(".", "").isdigit():
        obs = None  # a CONSIPA guardava número solto aqui (lixo do formulário)
    fazenda = "Castro" if setor_original == "PECUÁRIA DE LEITE" or "castro" in normalizar(obs) else "Água do Tigre"

    motivos = []
    central, categoria = _txt(r["central"]), _txt(r["categoria"])
    cod_cat = limpar_nome(categoria)[0] if categoria else None
    cod_cen = limpar_nome(central)[0] if central else None
    if cod_cat in CORRECOES_GRUPO and cod_cen and not cod_cat.startswith(cod_cen):
        central = CORRECOES_GRUPO[cod_cat]
        motivos.append("Categoria estava em outro grupo (corrigido na importação)")
    grupo = limpar_nome(central)[1] if central else None
    cat = limpar_nome(categoria)[1] if categoria else None
    if not cat:
        motivos.append("Sem categoria")
    if setor_original == "AGRICULTURA" and tipo == "despesa":
        motivos.append("Era 'Agricultura' na CONSIPA")
    if setor_original == "CONFIGURAÇÕES":
        motivos.append("Estava em 'Configurações' na CONSIPA")
    if cat and (tipo, setor, grupo, cat) not in plano:
        db.add(FinPlano(tipo=tipo, setor=setor, grupo=grupo, categoria=cat,
                        codigo=limpar_nome(categoria)[0] or limpar_nome(central)[0]))
        plano.add((tipo, setor, grupo, cat))

    nome_forn = _txt(r["fornecedor"])
    forn = None
    if nome_forn:
        forn = fornecedores.get(normalizar(nome_forn))
        if forn is None:
            forn = FinFornecedor(nome=nome_forn)
            db.add(forn)
            fornecedores[normalizar(nome_forn)] = forn

    item = _txt(r["item"])
    descricao = _txt(r["descricao"])
    if item and item not in (descricao or ""):
        descricao = f"{descricao} · {limpar_nome(item)[1]}" if descricao else limpar_nome(item)[1]

    lanc = FinLancamento(
        tipo=tipo, data=_dt(r["data"]), fazenda=fazenda, setor=setor, grupo=grupo, categoria=cat,
        fornecedor_nome=forn.nome if forn else None, descricao=descricao, conta=_txt(r["conta"]),
        nf_numero=_txt(r["nf"]), observacao=obs, origem="consipa", chave_origem=origem,
        criado_por=f"{usuario} (importação CONSIPA)", revisar=bool(motivos),
        revisar_motivo="; ".join(motivos) or None,
    )
    lanc.fornecedor = forn  # ligado abaixo (fornecedor pode ser novo, ainda sem id)
    for i, p in enumerate(bloco, start=1):
        lanc.parcelas.append(FinParcela(numero=i, vencimento=_dt(p["data_fluxo"]) or lanc.data,
                                        valor=round(float(p["valor"]), 2)))
    lanc.valor_total = round(sum(p.valor for p in lanc.parcelas), 2)
    return lanc


def _aprender_padroes(db: Session) -> None:
    """Cada fornecedor importado ganha como sugestão a classificação mais recente."""
    ult = {}
    q = (select(FinLancamento)
         .where(FinLancamento.categoria.is_not(None), FinLancamento.fornecedor_id.is_not(None))
         .order_by(FinLancamento.data))
    for lc in db.scalars(q):
        ult[lc.fornecedor_id] = lc
    for fid, lc in ult.items():
        f = db.get(FinFornecedor, fid)
        if f and not f.pad_categoria:
            f.pad_tipo, f.pad_fazenda, f.pad_setor = lc.tipo, lc.fazenda, lc.setor
            f.pad_grupo, f.pad_categoria, f.pad_conta = lc.grupo, lc.categoria, lc.conta
