"""Cria: vínculo mãe-bezerro, nascimento, desmama e situação das matrizes.

Regras (combinadas com o dono):
- Bezerro "ao pé" = cria ATIVA, ainda com tipo de bezerro e sem data de desmame.
- Matriz SOLTEIRA = não tem bezerro ao pé (nunca pariu, já desmamou ou a cria saiu).
- Prenhez é uma marcação opcional da vaca: "prenhe" (toque), "mojando" (visto no
  visual, perto de parir) ou "vazia" (toque). O nascimento de um bezerro limpa.
- Pode ir pro frigorífico = solteira e NÃO marcada como prenhe/mojando.
"""
from __future__ import annotations

from datetime import date

from sqlalchemy import func, or_
from sqlalchemy.orm import Session, selectinload

from ..models import Animal, AnimalLote, Pesagem, StatusAnimal
from .consultas import lote_atual

PRENHEZ_VALORES = ["prenhe", "mojando", "vazia"]
PRENHEZ_POSITIVA = {"prenhe", "mojando"}
TIPO_BEZERRO = {"F": "Bez Fem", "M": "Bez Mach"}
# Tipo que o bezerro vira ao desmamar (pelo sexo do tipo atual).
TIPO_DESMAMADO = {"F": "Novilha", "M": "Boi"}


def eh_bezerro(tipo: str | None) -> bool:
    return (tipo or "").strip().lower().startswith("bez")


def sexo_do_tipo(tipo: str | None) -> str | None:
    """'F'/'M' a partir do tipo do bezerro ("Bez Fem" / "Bez Mach")."""
    t = (tipo or "").lower()
    if "fem" in t:
        return "F"
    if "mach" in t:
        return "M"
    return None


def ao_pe(cria: Animal) -> bool:
    return (cria.status == StatusAnimal.ATIVO and cria.data_desmame is None
            and eh_bezerro(cria.tipo))


def _ultimo_peso(a: Animal) -> float | None:
    return a.pesagens[-1].peso if a.pesagens else None


def _cria_dict(c: Animal, hoje: date) -> dict:
    desmama = None
    if c.data_desmame:
        # Peso à desmama = pesagem mais próxima da data (até 15 dias de diferença).
        perto = [p for p in c.pesagens if abs((p.data - c.data_desmame).days) <= 15]
        if perto:
            desmama = min(perto, key=lambda p: abs((p.data - c.data_desmame).days)).peso
    return {
        "id": c.id, "brinco": c.brinco, "sem_brinco": c.sem_brinco, "tipo": c.tipo,
        "status": c.status.value, "nascimento": c.nascimento,
        "idade_dias": (hoje - c.nascimento).days if c.nascimento else None,
        "ultimo_peso": _ultimo_peso(c), "data_desmame": c.data_desmame,
        "peso_desmame": desmama, "ao_pe": ao_pe(c),
    }


def resumo_matriz(m: Animal, hoje: date | None = None) -> dict:
    """Situação de uma matriz + as crias dela (usado na lista e na ficha)."""
    hoje = hoje or date.today()
    crias = sorted(m.crias, key=lambda c: (c.nascimento or date.min, c.id))
    no_pe = [c for c in crias if ao_pe(c)]
    partos = sorted(c.nascimento for c in crias if c.nascimento)
    # Intervalo entre partos (média, em dias) quando há 2+ nascimentos com data.
    intervalos = [(b - a).days for a, b in zip(partos, partos[1:]) if (b - a).days > 0]
    solteira = not no_pe
    return {
        "id": m.id, "brinco": m.brinco, "tipo": m.tipo, "raca": m.raca,
        "lote": lote_atual(m), "ultimo_peso": _ultimo_peso(m),
        "situacao": "solteira" if solteira else "com_bezerro",
        "prenhez": m.prenhez, "prenhez_data": m.prenhez_data,
        "pode_frigorifico": solteira and m.prenhez not in PRENHEZ_POSITIVA,
        "crias": [_cria_dict(c, hoje) for c in crias],
        "total_crias": len(crias),
        "ultimo_parto": partos[-1] if partos else None,
        "intervalo_partos_dias": round(sum(intervalos) / len(intervalos)) if intervalos else None,
    }


def painel(db: Session, novilhas: bool = False) -> dict:
    """Tudo que a aba Cria mostra, em poucas consultas (o banco fica longe).

    Matrizes = vacas + qualquer animal que já tem cria; com `novilhas=True` entram
    também as novilhas (pra lançar a 1ª cria de uma novilha). Bezerros = TODOS os
    bezerros ativos, com ou sem mãe informada.
    """
    ids_maes = db.query(Animal.mae_id).filter(Animal.mae_id.isnot(None)).distinct()
    tipos = [func.lower(Animal.tipo).like("vaca%"), Animal.id.in_(ids_maes)]
    if novilhas:
        tipos.append(func.lower(Animal.tipo).like("novilha%"))
    matrizes = (
        db.query(Animal)
        .filter(Animal.status == StatusAnimal.ATIVO, or_(*tipos))
        .options(
            selectinload(Animal.pesagens),
            selectinload(Animal.lotes).selectinload(AnimalLote.lote),
            selectinload(Animal.crias).selectinload(Animal.pesagens),
        )
        .all()
    )
    hoje = date.today()
    lista = [resumo_matriz(m, hoje) for m in matrizes]

    def chave(b: str):
        b = (b or "").strip()
        return (0, int(b), "") if b.isdigit() else (1, 0, b.upper())
    lista.sort(key=lambda m: chave(m["brinco"]))

    bezerros = (
        db.query(Animal)
        .filter(Animal.status == StatusAnimal.ATIVO, func.lower(Animal.tipo).like("bez%"))
        .options(selectinload(Animal.pesagens),
                 selectinload(Animal.lotes).selectinload(AnimalLote.lote),
                 selectinload(Animal.mae))
        .all()
    )
    lista_bez = sorted(
        ({"id": b.id, "brinco": b.brinco, "sem_brinco": b.sem_brinco, "tipo": b.tipo,
          "sexo": sexo_do_tipo(b.tipo),
          "mae": {"id": b.mae.id, "brinco": b.mae.brinco} if b.mae else None,
          "nascimento": b.nascimento,
          "idade_dias": (hoje - b.nascimento).days if b.nascimento else None,
          "desmamado": b.data_desmame is not None,
          "lote": lote_atual(b), "ultimo_peso": _ultimo_peso(b)} for b in bezerros),
        key=lambda b: chave(b["brinco"]))
    return {
        "matrizes": lista,
        "bezerros": lista_bez,
        "resumo": {
            "matrizes": len(lista),
            "com_bezerro": sum(1 for m in lista if m["situacao"] == "com_bezerro"),
            "solteiras": sum(1 for m in lista if m["situacao"] == "solteira"),
            "prenhes": sum(1 for m in lista if m["prenhez"] in PRENHEZ_POSITIVA),
            "pode_frigorifico": sum(1 for m in lista if m["pode_frigorifico"]),
            "bezerros_ao_pe": sum(1 for m in lista for c in m["crias"] if c["ao_pe"]),
            "bezerros": len(lista_bez),
            "bezerros_sem_mae": sum(1 for b in lista_bez if b["mae"] is None),
            "bezerros_sem_nascimento": sum(1 for b in lista_bez if b["nascimento"] is None),
        },
    }


def gerar_excel(titulo: str, cabecalho: list[str], linhas: list[list]) -> bytes:
    """Excel do que está na tela da Cria (a tela manda as linhas já filtradas/ordenadas)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = (titulo or "Cria")[:31]
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append(linha)
    for i, nome in enumerate(cabecalho, start=1):
        maior = max([len(str(nome))] + [len(str(l[i - 1])) for l in linhas if i <= len(l)
                                         and l[i - 1] is not None])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(maior + 2, 40)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def registrar_nascimento(db: Session, mae: Animal, data: date, sexo: str,
                         brinco: str | None = None, peso: float | None = None) -> Animal:
    """Cadastra o bezerro já ligado à mãe, no mesmo lote dela."""
    sexo = (sexo or "").strip().upper()[:1]
    if sexo not in TIPO_BEZERRO:
        raise ValueError("Informe o sexo do bezerro (fêmea ou macho).")
    if data > date.today():
        raise ValueError("A data de nascimento não pode ser no futuro.")
    brinco = (brinco or "").strip()
    bezerro = Animal(
        brinco=brinco or "S/B", sem_brinco=not brinco, tipo=TIPO_BEZERRO[sexo],
        nascimento=data, mae_id=mae.id, status=StatusAnimal.ATIVO,
    )
    db.add(bezerro)
    db.flush()
    if not brinco:
        # Brinco provisório legível até o bezerro ser brincado: S/B-<mãe>-<id>.
        bezerro.brinco = f"S/B-{mae.brinco}-{bezerro.id}"
    atual = next((al for al in mae.lotes if al.data_fim is None), None)
    if atual is not None:
        db.add(AnimalLote(animal_id=bezerro.id, lote_id=atual.lote_id, data_inicio=data))
    if peso:
        db.add(Pesagem(animal_id=bezerro.id, data=data, peso=peso))
    # Pariu: a marcação de prenhe/mojando deixa de valer.
    mae.prenhez = None
    mae.prenhez_data = None
    db.commit()
    db.refresh(bezerro)
    return bezerro


def marcar_prenhez(db: Session, animal: Animal, prenhez: str | None,
                   data: date | None = None) -> None:
    prenhez = (prenhez or "").strip().lower() or None
    if prenhez is not None and prenhez not in PRENHEZ_VALORES:
        raise ValueError("Marcação inválida (use prenhe, mojando ou vazia).")
    animal.prenhez = prenhez
    animal.prenhez_data = (data or date.today()) if prenhez else None
    db.commit()


def definir_nascimento(animal: Animal, nascimento: date | None) -> None:
    if nascimento is not None and nascimento > date.today():
        raise ValueError("A data de nascimento não pode ser no futuro.")
    animal.nascimento = nascimento


def definir_mae(db: Session, cria: Animal, mae: Animal | None) -> None:
    if mae is not None:
        if mae.id == cria.id:
            raise ValueError("O animal não pode ser mãe dele mesmo.")
        if mae.mae_id == cria.id:
            raise ValueError("Esse animal é cria do bezerro informado — confira os brincos.")
    cria.mae_id = mae.id if mae else None
    db.commit()


def desmamar(db: Session, cria: Animal, data: date | None = None,
             novo_tipo: str | None = None) -> None:
    """Marca a desmama; o bezerro vira Novilha/Boi (ou o tipo informado)."""
    cria.data_desmame = data or date.today()
    if novo_tipo:
        cria.tipo = novo_tipo.strip()
    elif eh_bezerro(cria.tipo):
        cria.tipo = TIPO_DESMAMADO.get(sexo_do_tipo(cria.tipo) or "", cria.tipo)
    db.commit()
