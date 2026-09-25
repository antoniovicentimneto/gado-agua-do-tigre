"""Testes do relatório filtrado em Excel (Rebanho › Planilha)."""
import io

from openpyxl import load_workbook

from app.models import Animal
from app.services.relatorio import gerar_relatorio


def _linhas(conteudo):
    ws = load_workbook(io.BytesIO(conteudo)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_relatorio_respeita_ordem_e_filtro(db):
    ids = {a.brinco: a.id for a in db.query(Animal).all()}
    # Só dois animais, na ordem pedida (do mais pesado pro mais leve).
    linhas = _linhas(gerar_relatorio(db, [ids["201"], ids["101"]]))
    assert linhas[0][0] == "Brinco"
    assert [linhas[1][0], linhas[2][0]] == ["201", "101"]
    assert linhas[1][5] == 450 and linhas[2][5] == 400
    # Linha de totais: 2 animais, 850 kg.
    totais = linhas[-1]
    assert totais[1] == 2 and totais[4] == 850


def test_relatorio_vazio(db):
    linhas = _linhas(gerar_relatorio(db, []))
    assert linhas[0][0] == "Brinco" and linhas[-1][1] == 0
