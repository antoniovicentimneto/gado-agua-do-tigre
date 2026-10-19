"""Testes do Financeiro: lançamentos, parcelas, resultado por mês, recorrentes,
fornecedores, importação da CONSIPA e permissões da API."""
import io
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.models import PapelUsuario, Usuario
from app.models_financeiro import FinFornecedor, FinLancamento, FinPlano
from app.services import fin_consipa, fin_lancamentos as svc, fin_plano, fin_resultado
from app.services.auth import criar_token
from app.services.fin_plano import ErroFinanceiro


@pytest.fixture
def fdb():
    """Banco em memória só com as tabelas, com cadastros padrão do Financeiro."""
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    fin_plano.garantir_cadastros(s)
    yield s
    s.close()


def _lanc(db, **extra):
    dados = {"tipo": "despesa", "data": "2026-03-05", "fazenda": "Água do Tigre", "setor": "Sede",
             "grupo": "Gastos gerais", "categoria": "Mercado - gastos da fazenda",
             "fornecedor_nome": "Mercado Bom Preço", "descricao": "compras", "valor_total": 100,
             "conta": "Sicoob"}
    dados.update(extra)
    return svc.criar(db, dados, "Letícia")


# ---------------------------------------------------------------- plano / nomes

def test_limpar_nome_tira_codigo_e_mantem_siglas():
    assert fin_plano.limpar_nome("1.01.09. PRÓ LABORE") == ("1.01.09", "Pró labore")
    assert fin_plano.limpar_nome("1.05.10. ITR - IMPOSTO TERRITORIAL RURAL")[1] == "ITR - imposto territorial rural"
    assert fin_plano.limpar_nome("1.05.07. IPVA, LICENCIAMENTO")[1].startswith("IPVA,")


def test_cadastros_padrao_tem_castro_e_plano(fdb):
    cad = fin_plano.cadastros(fdb)
    assert [f["nome"] for f in cad["fazendas"]] == ["Água do Tigre", "Castro"]
    setores = {p["setor"] for p in cad["plano"]}
    assert "Agricultura" not in setores and "Pecuária" in setores


def test_renomear_categoria_atualiza_lancamentos(fdb):
    _lanc(fdb)
    q = select(FinPlano).where(FinPlano.setor == "Sede", FinPlano.categoria == "Mercado - gastos da fazenda")
    p = fdb.scalar(q)
    fin_plano.salvar_plano(fdb, {"tipo": "despesa", "setor": "Sede", "grupo": "Gastos gerais",
                                 "categoria": "Mercado"}, p.id)
    assert fdb.scalar(select(FinLancamento.categoria)) == "Mercado"


# ---------------------------------------------------------------- lançamentos e parcelas

def test_parcelas_dividem_centavos_na_ultima(fdb):
    lanc = _lanc(fdb, valor_total=100, n_parcelas=3, primeiro_vencimento="2026-01-31")
    assert [p.valor for p in lanc.parcelas] == [33.33, 33.33, 33.34]
    assert [p.vencimento for p in lanc.parcelas] == [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)]


def test_parcelas_que_nao_batem_com_total_sao_recusadas(fdb):
    with pytest.raises(ErroFinanceiro):
        _lanc(fdb, valor_total=100, parcelas=[{"vencimento": "2026-03-10", "valor": 50}])


def test_mudar_total_refaz_parcelas_mantendo_quantidade(fdb):
    lanc = _lanc(fdb, valor_total=300, n_parcelas=3)
    lanc = svc.atualizar(fdb, lanc.id, {"valor_total": 600}, "Antonio")
    assert [p.valor for p in lanc.parcelas] == [200, 200, 200]
    assert lanc.alterado_por == "Antonio"


def test_sem_categoria_vai_pra_revisao_e_sai_quando_conferido(fdb):
    lanc = _lanc(fdb, categoria="")
    assert lanc.revisar and lanc.revisar_motivo == "Sem categoria"
    lanc = svc.atualizar(fdb, lanc.id, {"categoria": "Mercado - gastos da fazenda", "conferido": True}, "Letícia")
    assert not lanc.revisar


def test_excluir_nao_apaga_e_some_da_lista(fdb):
    lanc = _lanc(fdb)
    svc.excluir(fdb, lanc.id, "lançado em dobro", "Antonio")
    assert svc.listar(fdb) == []
    assert svc.listar(fdb, excluidos=True)[0]["excluido_motivo"] == "lançado em dobro"
    svc.restaurar(fdb, lanc.id, "Antonio")
    assert len(svc.listar(fdb)) == 1


def test_valor_em_formato_brasileiro(fdb):
    lanc = _lanc(fdb, valor_total="1.234,56")
    assert lanc.valor_total == 1234.56


# ---------------------------------------------------------------- fornecedores

def test_fornecedor_acha_sem_acento_e_aprende_classificacao(fdb):
    _lanc(fdb, fornecedor_nome="Frísia")
    f = svc.achar_ou_criar_fornecedor(fdb, "FRISIA", criar=False)
    assert f is not None and f.nome == "Frísia"
    sug = svc.sugestao_fornecedor(fdb, "frisia")
    assert sug["categoria"] == "Mercado - gastos da fazenda" and sug["conta"] == "Sicoob"


def test_juntar_fornecedores_move_lancamentos_e_guarda_apelido(fdb):
    a = _lanc(fdb, fornecedor_nome="Piscinas e CIA")
    b = _lanc(fdb, fornecedor_nome="Piscina e cia", valor_total=50)
    svc.juntar_fornecedores(fdb, a.fornecedor_id, [b.fornecedor_id])
    fdb.refresh(b)
    assert b.fornecedor_id == a.fornecedor_id
    f = fdb.get(FinFornecedor, a.fornecedor_id)
    assert "Piscina e cia" in f.apelidos
    assert svc.achar_ou_criar_fornecedor(fdb, "piscina e cia", criar=False).id == f.id


def test_duplicado_mesmo_fornecedor_valor_e_data_proxima(fdb):
    a = _lanc(fdb, data="2026-03-05")
    b = _lanc(fdb, data="2026-03-07")
    _lanc(fdb, data="2026-04-20")
    assert svc.possiveis_duplicados(fdb) == [[a.id, b.id]]
    svc.marcar_nao_duplicado(fdb, [a.id, b.id])
    assert svc.possiveis_duplicados(fdb) == []


def test_parcelas_com_descricao_diferente_nao_sao_duplicado(fdb):
    a = _lanc(fdb, data="2026-03-05", descricao="Motobomba VT 1786,30 1/2P")
    _lanc(fdb, data="2026-03-05", descricao="Motobomba VT 1786,30 2/2P")
    c = _lanc(fdb, data="2026-03-06", descricao="motobomba  VT 1786,30 1/2P")  # só espaço/maiúscula
    assert svc.possiveis_duplicados(fdb) == [[a.id, c.id]]


# ---------------------------------------------------------------- resultado por mês

def test_resultado_caixa_por_mes_e_detalhe(fdb):
    _lanc(fdb, valor_total=300, n_parcelas=3, primeiro_vencimento="2026-03-10")
    svc.criar(fdb, {"tipo": "receita", "data": "2026-03-01", "setor": "Sede", "grupo": "Arrendamento",
                    "categoria": "Soja", "valor_total": 1000}, "Antonio")
    res = fin_resultado.resultado(fdb, {"inicio": "2026-03", "fim": "2026-04"})
    assert res["meses"] == ["2026-03", "2026-04"]
    desp = res["arvore"][0]
    assert desp["nome"] == "despesa" and desp["meses"] == {"2026-03": 100, "2026-04": 100}
    assert res["saldo"] == {"2026-03": 900, "2026-04": -100}
    linhas = fin_resultado.detalhar(fdb, {"inicio": "2026-03", "fim": "2026-04"},
                                    ["despesa", "Sede", "Gastos gerais"], "2026-04")
    assert len(linhas) == 1 and linhas[0]["parcela"] == "2/3" and linhas[0]["valor"] == 100


def test_resultado_competencia_usa_mes_da_compra(fdb):
    _lanc(fdb, valor_total=300, n_parcelas=3, primeiro_vencimento="2026-03-10")
    res = fin_resultado.resultado(fdb, {"visao": "competencia"})
    assert res["arvore"][0]["meses"] == {"2026-03": 300}


def test_resultado_separa_fazenda(fdb):
    _lanc(fdb, fazenda="Castro", valor_total=70)
    _lanc(fdb, valor_total=30, fornecedor_nome="Outro")
    res = fin_resultado.resultado(fdb, {"fazenda": "Castro"})
    assert res["arvore"][0]["total"] == 70


# ---------------------------------------------------------------- recorrentes

def test_recorrente_gera_todo_mes_sem_duplicar(fdb):
    svc.salvar_recorrente(fdb, {"descricao": "Salário Romário", "setor": "Pecuária", "grupo": "Recursos humanos",
                                "categoria": "Salário", "valor": 2500, "dia": 5, "inicio": "2026-01-01"}, None)
    hoje = date(2026, 3, 20)
    svc.gerar_recorrentes(fdb, hoje)
    assert svc.gerar_recorrentes(fdb, hoje) == 0
    meses = sorted(fdb.scalars(select(FinLancamento.recorrente_mes)))
    assert len(meses) == len(set(meses))  # um por mês, nunca dois
    assert meses[:4] == ["2026-01", "2026-02", "2026-03", "2026-04"]
    lanc = fdb.scalar(select(FinLancamento).where(FinLancamento.recorrente_mes == "2026-02"))
    assert lanc.parcelas[0].pago is False and lanc.parcelas[0].vencimento == date(2026, 2, 5)


# ---------------------------------------------------------------- importação da CONSIPA

def _planilha_consipa() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CentralReal"
    ws.append([])
    ws.append([])
    ws.append([None] + [c.upper() for c in ["DATA CONTÁBIL"] + fin_consipa.COLUNAS[1:]])
    linha = lambda **k: [None] + [k.get(c) for c in fin_consipa.COLUNAS]  # noqa: E731
    base = dict(data=date(2025, 5, 2), status="REALIZADO", es="1. DESPESA", nf=123, descricao="sal mineral",
                n_parcelas=2, qtde=1, valor_unit=1000, conta="Sicoob")
    ws.append(linha(**base, data_fluxo=date(2025, 5, 10), setor="PECUÁRIA DE LEITE",
                    central="1.02. ALIMENTOS", categoria="1.02.16. MINERAIS", fornecedor="Frisia", valor=500))
    ws.append(linha(**base, data_fluxo=date(2025, 6, 10), setor="PECUÁRIA DE LEITE",
                    central="1.02. ALIMENTOS", categoria="1.02.16. MINERAIS", fornecedor="Frisia", valor=500))
    ws.append(linha(data=date(2025, 7, 1), data_fluxo=date(2025, 7, 1), es="1. DESPESA", setor="PESSOAL",
                    central="1.12. RETIRADA DE LUCRO", categoria="1.04.02. LUZ ELÉTRICA", fornecedor="Copel",
                    descricao="luz", n_parcelas=1, valor=80))
    ws.append(linha(data=date(2025, 7, 2), data_fluxo=date(2025, 7, 2), es="1. DESPESA", setor="SEDE",
                    central="1.10. GASTOS GERAIS", categoria=0, fornecedor="Frísia", descricao="diversos",
                    n_parcelas=1, valor=40))
    ws.append(linha(data=date(2025, 4, 1), es=" 5. ATIVIDADE ", setor="INVENTÁRIO", valor=999))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_importar_consipa_agrupa_parcelas_e_aplica_decisoes(fdb):
    r = fin_consipa.importar(fdb, _planilha_consipa(), "Antonio")
    assert r["lancamentos"] == 3 and r["parcelas"] == 4 and r["linhas_ignoradas"] == 1
    sal = fdb.scalar(select(FinLancamento).where(FinLancamento.descricao == "sal mineral"))
    assert sal.fazenda == "Castro" and sal.setor == "Pecuária" and sal.valor_total == 1000
    assert [p.valor for p in sal.parcelas] == [500, 500]
    luz = fdb.scalar(select(FinLancamento).where(FinLancamento.descricao == "luz"))
    assert luz.grupo == "Materiais de consumo, operação e manutenção" and luz.revisar
    div = fdb.scalar(select(FinLancamento).where(FinLancamento.descricao == "diversos"))
    assert div.categoria is None and div.revisar_motivo == "Sem categoria"
    # Frisia e Frísia viram o mesmo fornecedor
    assert sal.fornecedor_id == div.fornecedor_id
    # Rodar de novo não duplica
    assert fin_consipa.importar(fdb, _planilha_consipa(), "Antonio")["lancamentos"] == 0


def _planilha_antiga() -> bytes:
    """Versão antiga: sem a coluna FAZENDA e com o VALOR TOTAL quebrado (#REF!), em duas abas."""
    wb = openpyxl.Workbook()
    cab = ["DATA CONTÁBIL", "DATA FLUXO", "STATUS", "SETOR", "ENTRADAS/SAÍDAS", "CENTRAL DE CUSTO", "CATEGORIA",
           "ITEM", "SUBITEM", "Nº NOTA FISCAL", "FORNECEDOR", "DESCRIÇÃO", "Nº DE PARCELAS", "QUANTIDADE",
           "VALOR UNITÁRIO", "@/KG/L/SACA", "UNIDADE", "RENDIMENTO", "DESCONTO / ÁGIO", "OBSERVAÇÃO", "CONTA",
           "VALOR TOTAL"]
    ws = wb.active
    ws.title = "Planilha1"
    ws.append([None] + cab)

    def linha(d, desc, q, vu, arroba=0, desc_pct=0, parc=1, setor="SEDE"):
        return [None, d, d, "REALIZADO", setor, "1. DESPESA", "1.10. GASTOS GERAIS", "1.10.02. MERCADO", None, None,
                None, "Loja", desc, parc, q, vu, arroba, None, None, desc_pct, None, None, "#REF!"]
    ws.append(linha(date(2024, 3, 1), "compra antiga", 2, 100, desc_pct=0.1))         # 2×100×0,9 = 180
    ws.append(linha(date(2024, 4, 1), "boi por kg", 10, 9.5, arroba=400, parc=2))   # 10×9,5×400÷2 = 19.000 (1 parcela)
    ws.append(linha(date(2025, 5, 2), "depois do corte", 1, 50))                    # vale a CONSIPA
    ws2 = wb.create_sheet("Planilha2")
    ws2.append([None] + cab)
    ws2.append(linha(date(2024, 3, 1), "compra antiga", 2, 100, desc_pct=0.1))       # cópia: não repete
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_valor_pela_formula_da_consipa():
    f = fin_consipa.valor_pela_formula
    assert f({"qtde": 580, "valor_unit": 140.5, "desconto": 0.015}) == 80267.65
    assert f({"qtde": 19, "valor_unit": 316.5988, "arroba": 18.424, "desconto": 0.015}) == pytest.approx(109164.9, abs=0.01)
    assert f({"qtde": 1, "valor_unit": 22450, "desconto": 0.13367, "n_parcelas": 5}) == pytest.approx(3889.82, abs=0.01)


def test_importar_com_planilha_antiga(fdb):
    r = fin_consipa.importar(fdb, _planilha_consipa(), "Antonio", _planilha_antiga())
    info = r["planilha_antiga"]
    assert info["antiga_incluidas"] == 2 and info["antiga_repetidas"] == 1  # a cópia da Planilha2
    antiga = fdb.scalar(select(FinLancamento).where(FinLancamento.descricao.like("compra antiga%")))
    assert antiga.valor_total == 180
    boi = fdb.scalar(select(FinLancamento).where(FinLancamento.descricao.like("boi por kg%")))
    assert boi.valor_total == 19000  # 10 × 9,5 × 400 ÷ 2: uma parcela (cada parcela é uma linha)
    assert fdb.scalar(select(FinLancamento).where(FinLancamento.descricao.like("depois do corte%"))) is None
    # De novo: nada novo
    assert fin_consipa.importar(fdb, _planilha_consipa(), "Antonio", _planilha_antiga())["lancamentos"] == 0


def test_texto_comprido_cabe_na_coluna():
    lc = fin_consipa._caber(FinLancamento(tipo="despesa", nf_numero="IMPOSTO / ENCARGOS DA FOLHA", conta="x" * 200))
    assert lc.nf_numero == "IMPOSTO / ENCARGOS D" and len(lc.conta) == 80


def test_importar_arquivo_que_nao_e_consipa(fdb):
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ErroFinanceiro):
        fin_consipa.importar(fdb, buf.getvalue(), "Antonio")


# ---------------------------------------------------------------- API e permissões

@pytest.fixture
def cliente():
    from app.main import app
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    Sessao = sessionmaker(bind=engine)
    s = Sessao()
    dono = Usuario(nome="Antonio", usuario="antonio", senha_hash="x", papel=PapelUsuario.DONO)
    peao = Usuario(nome="Romário", usuario="romario", senha_hash="x", papel=PapelUsuario.PEAO)
    s.add_all([dono, peao])
    s.commit()
    tokens = {"dono": criar_token(dono), "peao": criar_token(peao)}
    s.close()

    def _db():
        d = Sessao()
        try:
            yield d
        finally:
            d.close()
    app.dependency_overrides[get_db] = _db
    yield TestClient(app), tokens
    app.dependency_overrides.clear()


def test_peao_nao_ve_financeiro(cliente):
    c, tokens = cliente
    r = c.get("/api/fin/cadastros", headers={"Authorization": f"Bearer {tokens['peao']}"})
    assert r.status_code == 403


def test_api_lancar_e_ver_resultado(cliente):
    c, tokens = cliente
    h = {"Authorization": f"Bearer {tokens['dono']}"}
    assert c.get("/api/fin/cadastros", headers=h).status_code == 200
    r = c.post("/api/fin/lancamentos", headers=h, json={
        "tipo": "despesa", "data": "2026-05-02", "setor": "Sede", "grupo": "Gastos gerais",
        "categoria": "Outros", "fornecedor_nome": "Loja", "valor_total": 90, "n_parcelas": 3})
    assert r.status_code == 200 and len(r.json()["parcelas"]) == 3
    r = c.post("/api/fin/lancamentos", headers=h, json={"tipo": "despesa", "valor_total": 0})
    assert r.status_code == 400 and "valor" in r.json()["detail"]
    res = c.get("/api/fin/resultado?inicio=2026-05&fim=2026-07", headers=h).json()
    assert res["arvore"][0]["total"] == 90


def test_api_cpf_invalido_recusado(cliente):
    c, tokens = cliente
    h = {"Authorization": f"Bearer {tokens['dono']}"}
    r = c.post("/api/fin/contribuintes", headers=h, json={"documento": "111.111.111-11", "nome": "X"})
    assert r.status_code == 400
    r = c.post("/api/fin/contribuintes", headers=h, json={"documento": "529.982.247-25", "nome": "Produtor"})
    assert r.status_code == 200 and r.json()["documento"] == "52998224725"
    r = c.post("/api/fin/contribuintes", headers=h, json={"documento": "11.222.333/0001-81", "nome": "Empresa"})
    assert r.status_code == 200
